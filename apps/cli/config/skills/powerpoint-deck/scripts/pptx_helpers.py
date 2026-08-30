"""pptx_helpers — low-level primitives for hand-designed python-pptx decks.

These helpers encode only the fiddly OOXML/python-pptx knowledge (bullets,
letter-spacing, shadows, native-chart styling, cell borders, image crop).
They take YOUR colors and impose NO layout, palette, or slide structure —
the design of each deck is yours to invent. See references/design.md in the
skill for palette and layout guidance.

    import sys; sys.path.insert(0, "<skill-dir>/scripts")
    from pptx_helpers import style_chart, bullet, soft_shadow, ...
"""

from __future__ import annotations

import math

from PIL import Image
from pptx.chart.chart import Chart
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

__all__ = [
    "rgb", "blend", "lighten", "darken", "tints",
    "wrap_lines", "no_autofit", "bullet", "letter_spacing",
    "strip_style", "soft_shadow", "cell_border_bottom", "crop_to",
    "style_chart",
]


# ----------------------------------------------------------- color math ---

def rgb(hex_color: str) -> RGBColor:
    """RGBColor from a bare hex string, with or without '#'.
    Python-pptx never accepts '#' inside color strings passed elsewhere —
    always route literals through this."""
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def blend(a: str, b: str, t: float) -> str:
    """Hex color ``a`` mixed toward ``b`` by ``t`` (0..1)."""
    a, b = a.lstrip("#"), b.lstrip("#")
    ca = tuple(int(a[i : i + 2], 16) for i in (0, 2, 4))
    cb = tuple(int(b[i : i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(ca[i] + (cb[i] - ca[i]) * t):02X}" for i in range(3))


def lighten(hex_color: str, t: float) -> str:
    return blend(hex_color, "FFFFFF", t)


def darken(hex_color: str, t: float) -> str:
    return blend(hex_color, "000000", t)


def tints(hex_color: str, count: int, lo: float = 0.25, hi: float = 0.85) -> list[str]:
    """``count`` shades of one hue, dark → light — for series ramps that stay
    on-palette without inventing new hues."""
    return [blend(hex_color, "FFFFFF", lo + (hi - lo) * i / max(1, count - 1)) for i in range(count)]


# ------------------------------------------------------------ text math ---

def wrap_lines(text: str, width_in: float, font_pt: float, factor: float = 0.52) -> int:
    """Conservative wrapped-line estimate for mixed-case text (for sizing
    boxes before render). Bump ``factor`` for bold or serif faces."""
    if not text:
        return 1
    chars_per_line = max(8, int(width_in * 72 / (font_pt * factor)))
    return max(1, math.ceil(len(text) / chars_per_line))


def no_autofit(text_frame) -> None:
    """Word-wrap without resize — overflow then shows up in visual QA instead
    of silently shrinking."""
    text_frame.word_wrap = True
    text_frame.auto_size = None
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0


def bullet(paragraph, color_hex: str, char: str = "•", indent_in: float = 0.24) -> None:
    """Colored hanging-indent bullet. python-pptx has no bullet API; without
    this, wrapped bullet lines outdent to the margin."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = pPr.makeelement(qn("a:srgbClr"), {"val": color_hex.lstrip("#").upper()})
    buClr.append(srgb)
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


def letter_spacing(run, hundredths_pt: int) -> None:
    """Tracking. ~120–180 for uppercase kickers/eyebrows; body text stays 0."""
    run._r.get_or_add_rPr().set("spc", str(hundredths_pt))


# ---------------------------------------------------------------- shapes ---

def strip_style(shape) -> None:
    """Remove the theme style reference python-pptx adds to every autoshape.
    Without this PowerPoint may layer theme effects over your explicit
    fill/line, and LibreOffice renders the two differently."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def soft_shadow(shape, color_hex: str = "3A4560", alpha: float = 0.20,
                blur_pt: float = 7.0, dist_pt: float = 2.8) -> None:
    """Explicit subtle drop shadow (declared, so PowerPoint and LibreOffice
    agree). Also disables the inherited shadow first."""
    shape.shadow.inherit = False
    spPr = shape._element.spPr
    effect = spPr.makeelement(qn("a:effectLst"), {})
    shdw = spPr.makeelement(qn("a:outerShdw"), {
        "blurRad": str(int(blur_pt * 12700)), "dist": str(int(dist_pt * 12700)),
        "dir": "5400000", "rotWithShape": "0",
    })
    clr = spPr.makeelement(qn("a:srgbClr"), {"val": color_hex.lstrip("#").upper()})
    clr.append(spPr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))
    shdw.append(clr)
    effect.append(shdw)
    spPr.append(effect)


def cell_border_bottom(cell, color_hex: str, width_pt: float) -> None:
    """Hairline bottom rule for table cells (the only border a modern table
    needs). Insert at index 0 — tcPr schema puts line elements first."""
    tcPr = cell._tc.get_or_add_tcPr()
    ln = tcPr.makeelement(qn("a:lnB"), {"w": str(int(width_pt * 12700)), "cap": "flat"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    clr = ln.makeelement(qn("a:srgbClr"), {"val": color_hex.lstrip("#").upper()})
    fill.append(clr)
    ln.append(fill)
    tcPr.insert(0, ln)


def crop_to(path: str, aspect: float) -> str:
    """Center-crop an image to ``aspect`` (w/h) so half-bleed placements fill
    their frame without distortion. Returns the cropped file's path."""
    img = Image.open(path)
    w, h = img.size
    if w / h > aspect:
        new_w = int(h * aspect)
        box = ((w - new_w) // 2, 0, (w - new_w) // 2 + new_w, h)
    else:
        new_h = int(w / aspect)
        box = (0, (h - new_h) // 2, w, (h - new_h) // 2 + new_h)
    out = img.crop(box)
    cropped = path.rsplit(".", 1)[0] + f"-crop{out.size[0]}x{out.size[1]}.png"
    out.save(cropped)
    return cropped


# ---------------------------------------------------------------- charts ---

def style_chart(
    chart: Chart,
    colors: list[str],
    *,
    value_format: str | None = None,
    labels: bool | None = None,       # None = decide from shape (see below)
    gridlines: bool | None = None,    # None = decide from shape
    legend: bool | None = None,       # None = multi-series only
    hole: int = 62,                   # doughnut hole percent
) -> None:
    """Style a native chart with the deck's palette. One call, quiet frame:

    - no chart title (the slide title carries the claim)
    - data labels where they replace the value axis (columns/bars ≤16 points,
      single-series lines ≤8 points) — then axis and gridlines drop away
    - gridlines hairline-light when an axis remains
    - legend only for multi-series (bottom; right for pie/doughnut)
    - doughnut: per-slice colors darkened so the white inside labels read

    ``colors``: series colors in order (per-slice for pie/doughnut).
    """
    chart.has_title = False
    chart.font.size = Pt(10.5)
    chart.font.name = "Calibri"
    kind = "other"
    try:
        from pptx.enum.chart import XL_CHART_TYPE
        ct = chart.chart_type
        if ct in (XL_CHART_TYPE.PIE,):
            kind = "pie"
        elif ct in (XL_CHART_TYPE.DOUGHNUT,):
            kind = "doughnut"
        elif ct in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
            kind = "line"
        elif ct in (XL_CHART_TYPE.AREA, XL_CHART_TYPE.AREA_STACKED):
            kind = "area"
        elif ct in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.COLUMN_STACKED,
                    XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_STACKED):
            kind = "column"
    except Exception:
        pass

    plot = chart.plots[0]
    series = list(chart.series)
    n_series = len(series)
    try:
        n_cats = len(list(plot.categories))
    except Exception:
        n_cats = 0
    single = n_series == 1

    if kind in ("pie", "doughnut"):
        for i, pt in enumerate(series[0].points):
            base = colors[i % len(colors)]
            color = darken(base, 0.30) if kind == "doughnut" else base
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(color)
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_percentage = True
        dl.show_value = False
        dl.font.size = Pt(10.5)
        if kind == "pie":
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.font.color.rgb = rgb("333333")
        else:
            el = plot._element.find(qn("c:holeSize"))
            if el is not None:
                el.set("val", str(hole))
            dl.font.bold = True
            dl.font.color.rgb = rgb("FFFFFF")
        chart.has_legend = True if legend is None else legend
        _legend(chart, right=True)
        return

    # series colors + line styling
    for i, s in enumerate(series):
        color = colors[i % len(colors)]
        if kind == "line":
            s.format.line.color.rgb = rgb(color)
            s.format.line.width = Pt(2.4)
            s.smooth = False
            if getattr(chart.chart_type, "name", "").endswith("MARKERS") and 0 < n_cats <= 14:
                s.marker.style = XL_MARKER_STYLE.CIRCLE
                s.marker.size = 6
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = rgb(color)
                s.marker.format.line.fill.background()
        elif kind == "area":
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(lighten(color, 0.45))
            s.format.line.color.rgb = rgb(color)
            s.format.line.width = Pt(2.0)
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(color)

    # auto decisions when not forced
    if labels is None:
        labels = ((kind == "column" and n_cats * max(n_series, 1) <= 16)
                  or (kind == "line" and single and n_cats <= 8))
    if gridlines is None:
        gridlines = kind in ("line", "area") and not labels or n_cats > 8
    stacked = "STACKED" in getattr(chart.chart_type, "name", "")

    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.position = XL_LABEL_POSITION.CENTER if stacked else XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(10)
        dl.font.bold = stacked
        dl.font.color.rgb = rgb("333333")
        if value_format:
            dl.number_format = value_format
            dl.number_format_is_linked = False

    value_axis_visible = not labels
    va = chart.value_axis
    va.visible = value_axis_visible
    va.has_major_gridlines = bool(gridlines and value_axis_visible)
    if va.has_major_gridlines:
        gl = va.major_gridlines.format.line
        gl.color.rgb = rgb("E3E7EC")
        gl.width = Pt(0.75)
    if value_axis_visible:
        va.tick_labels.font.size = Pt(10)
        va.tick_labels.font.color.rgb = rgb("7A8494")
        if value_format:
            va.tick_labels.number_format = value_format
            va.tick_labels.number_format_is_linked = False
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    va.format.line.fill.background()

    if kind != "scatter":
        ca = chart.category_axis
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.minor_tick_mark = XL_TICK_MARK.NONE
        ca.format.line.color.rgb = rgb("D5DAE2")
        ca.format.line.width = Pt(1.0)
        ca.has_major_gridlines = False
        ca.tick_labels.font.size = Pt(10.5)
        ca.tick_labels.font.color.rgb = rgb("3A4354")

    chart.has_legend = (n_series > 1) if legend is None else legend
    if chart.has_legend:
        _legend(chart)

    if kind == "column" and not stacked:
        try:
            plot.gap_width = 55 if single else 70
            if not single:
                plot.overlap = -10
        except Exception:
            pass
    elif stacked:
        try:
            plot.gap_width = 70
            plot.overlap = 100
        except Exception:
            pass


def _legend(chart, right: bool = False) -> None:
    chart.legend.position = XL_LEGEND_POSITION.RIGHT if right else XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    chart.legend.font.color.rgb = rgb("7A8494")
